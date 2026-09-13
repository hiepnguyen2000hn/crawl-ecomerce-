"""Xuất dữ liệu đã crawl ra file Excel.

Giá trong DB lưu dạng **integer minor unit** (đúng quy ước, không sai số khi cộng dồn),
nhưng người đọc Excel cần số thực để lọc/sắp/tính. Nên ở đây quy đổi về đơn vị chính
(8999 → 89.99) và gắn định dạng số của Excel — cột vẫn là số, không phải chuỗi, nên
vẫn sort và dùng công thức được.
"""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.crawl.normalize import exponent_for

_HEADER_FILL = PatternFill("solid", fgColor="0A6C7E")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_MAX_WIDTH = 60


def minor_to_major(value: int | None, currency: str | None) -> float | None:
    if value is None:
        return None
    return value / (10 ** exponent_for(currency))


def _write_sheet(
    ws: Worksheet,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    *,
    money_cols: Sequence[int] = (),
    currency_col: int | None = None,
) -> None:
    ws.append(list(headers))
    for cell in ws[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    for row in rows:
        ws.append(list(row))

    # Định dạng tiền theo đúng mã tiền tệ của từng dòng — dataset trộn nhiều sàn
    # nên không thể gắn cứng một ký hiệu cho cả cột.
    if money_cols and currency_col:
        for r in range(2, ws.max_row + 1):
            cur = ws.cell(row=r, column=currency_col).value or ""
            fmt = f'#,##0.00 "{cur}"' if cur else "#,##0.00"
            for c in money_cols:
                ws.cell(row=r, column=c).number_format = fmt

    for idx, header in enumerate(headers, start=1):
        longest = len(str(header))
        for r in range(2, min(ws.max_row, 200) + 1):  # lấy mẫu 200 dòng là đủ
            v = ws.cell(row=r, column=idx).value
            if v is not None:
                longest = max(longest, len(str(v)))
        ws.column_dimensions[get_column_letter(idx)].width = min(longest + 2, _MAX_WIDTH)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def build_workbook(
    products: Sequence[Any],
    price_points: Sequence[Any],
    threads: Sequence[Any] = (),
    velocity: Sequence[dict] = (),
    ads: Sequence[Any] = (),
    landing_pages: Sequence[dict] = (),
) -> BytesIO:
    wb = Workbook()

    # ── Sản phẩm ─────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = "San pham"
    _write_sheet(
        ws,
        [
            "Nguon", "Domain", "Ma san pham", "Ten san pham", "Thuong hieu",
            "Tien te", "Gia thap nhat", "Gia cao nhat",
            "Danh gia", "So review", "So da ban", "Con hang",
            "Nguoi ban", "Quang cao",
            "Lan dau thay", "Lan cuoi thay", "URL",
        ],
        [
            (
                p.source, p.shop_domain, p.external_id, p.title, p.brand,
                p.currency,
                minor_to_major(p.price_min_minor, p.currency),
                minor_to_major(p.price_max_minor, p.currency),
                float(p.rating) if p.rating is not None else None,
                p.review_count,
                p.sales_volume,
                p.available,
                p.seller,
                p.is_sponsored,
                _naive(p.first_seen_at),
                _naive(p.last_seen_at),
                p.url,
            )
            for p in products
        ],
        money_cols=(7, 8),
        currency_col=6,
    )

    # ── Lịch sử giá ──────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Lich su gia")
    _write_sheet(
        ws2,
        [
            "Nguon", "Domain", "Ma san pham", "Ma bien the", "Ten bien the",
            "SKU", "Tien te", "Gia", "Gia gach ngang", "Con hang", "Ghi nhan luc",
        ],
        [
            (
                pp.source, pp.shop_domain, pp.external_product_id, pp.external_variant_id,
                pp.variant_title, pp.sku, pp.currency,
                minor_to_major(pp.price_minor, pp.currency),
                minor_to_major(pp.compare_at_minor, pp.currency),
                pp.available,
                _naive(pp.observed_at),
            )
            for pp in price_points
        ],
        money_cols=(8, 9),
        currency_col=7,
    )

    # ── Tín hiệu quảng cáo — SRS Bước 2.1 ───────────────────────────────────
    # `Ngay chay` (active_days) là bằng chứng mạnh nhất về việc sản phẩm có đang
    # sinh lời không: đó là tiền của chính đối thủ, không ai đốt ngân sách 60 ngày
    # cho một mẫu lỗ.
    if ads:
        wsa = wb.create_sheet("Tin hieu Ads")
        _write_sheet(
            wsa,
            [
                "Nha quang cao", "Ngay chay", "Dang chay", "Bat dau", "Ket thuc",
                "Quoc gia", "Media", "CTA", "So bien the",
                "Noi dung ads", "Landing page", "Tu khoa khop", "Link Ad Library",
            ],
            [
                (
                    a.page_name, a.active_days, a.is_active,
                    a.start_date, a.end_date, a.country,
                    a.media_type, a.cta_text, a.collation_count,
                    (a.ad_copy or "")[:500], a.landing_page_url,
                    a.matched_query, a.ad_library_url,
                )
                for a in ads
            ],
        )

    # ── Đối thủ nên soi giá — mắt xích Bước 2 sang Bước 3 ───────────────────
    if landing_pages:
        wslp = wb.create_sheet("Doi thu can soi gia")
        _write_sheet(
            wslp,
            ["Nha quang cao", "Quoc gia", "Ngay chay dai nhat", "So ads", "Landing page"],
            [
                (l["page_name"], l["country"], l["max_active_days"], l["ad_count"],
                 l["landing_page_url"])
                for l in landing_pages
            ],
        )

    # ── Tốc độ bán (proxy thay cho "số đã bán") ─────────────────────────────
    # Sheet này mới là thứ trả lời "sản phẩm nào bán chạy" — xem docstring
    # app/models/tracking.py. Cột "Du lieu du tin cay" là cố ý: con số tính trên
    # 2 ngày rất khác con số tính trên 30 ngày, người đọc phải thấy được.
    if velocity:
        ws4 = wb.create_sheet("Toc do ban")
        _write_sheet(
            ws4,
            [
                "Nguon", "Ten san pham", "Review hien tai", "Review tang them",
                "So ngay theo doi", "Review moi ngay", "Du lieu du tin cay",
                "Danh gia", "Nguoi ban", "Quang cao", "URL",
            ],
            [
                (
                    v["source"], v["title"], v["review_hien_tai"], v["review_tang_them"],
                    v["so_ngay_theo_doi"], v["review_moi_ngay"], v["du_lieu_du_tin_cay"],
                    v["rating"], v["seller"], v["is_sponsored"], v["url"],
                )
                for v in velocity
            ],
        )

    # ── Reddit VOC (chỉ tạo sheet khi có dữ liệu) ───────────────────────────
    if threads:
        ws3 = wb.create_sheet("Reddit VOC")
        _write_sheet(
            ws3,
            [
                "Subreddit", "Tieu de", "Diem", "So binh luan",
                "Tu khoa khop", "Dang luc", "Da cao binh luan", "URL",
            ],
            [
                (
                    t.subreddit, t.title, t.score, t.num_comments,
                    ", ".join(t.matched_keywords or []),
                    _naive(t.posted_at), t.comments_fetched, t.url,
                )
                for t in threads
            ],
        )

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _naive(dt: datetime | None) -> datetime | None:
    """Excel không hiểu timezone — openpyxl ném lỗi với datetime có tzinfo.
    Quy về UTC rồi bỏ tzinfo (dữ liệu vốn đã lưu UTC)."""
    if dt is None:
        return None
    return dt.astimezone(timezone.utc).replace(tzinfo=None)


def filename(prefix: str = "crawl-data") -> str:
    return f"{prefix}-{datetime.now(timezone.utc):%Y%m%d-%H%M}.xlsx"
